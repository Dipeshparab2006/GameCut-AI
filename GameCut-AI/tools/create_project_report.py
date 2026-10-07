from pathlib import Path
from datetime import date
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'output' / 'report'
IMG = OUT / 'figures'
OUT.mkdir(parents=True, exist_ok=True); IMG.mkdir(parents=True, exist_ok=True)

NAVY = '17324D'; BLUE = '2F75B5'; TEAL = '2A9D8F'; LIGHT = 'EAF3F8'; GREY = '5B6573'

def font(size=24, bold=False):
    try: return ImageFont.truetype('C:/Windows/Fonts/arial.ttf', size=size)
    except: return ImageFont.load_default()

def diagram(name, title, boxes, arrows, width=1500, height=850):
    im=Image.new('RGB',(width,height),'white'); d=ImageDraw.Draw(im)
    d.rectangle((0,0,width,92), fill='#17324D'); d.text((45,28),title,fill='white',font=font(34,True))
    coords={}
    for b in boxes:
        x,y,w,h,label=b; coords[label]=(x,y,w,h)
        d.rounded_rectangle((x,y,x+w,y+h),radius=20,fill='#EAF3F8',outline='#2F75B5',width=5)
        lines=label.split('\n'); yy=y+22
        for line in lines:
            bb=d.textbbox((0,0),line,font=font(25,True)); d.text((x+(w-(bb[2]-bb[0]))/2,yy),line,fill='#17324D',font=font(25,True)); yy+=35
    for a,b in arrows:
        x,y,w,h=coords[a]; x2,y2,w2,h2=coords[b]
        p1=(x+w,y+h//2); p2=(x2,y2+h2//2) if x2>x+w else (x2+w2//2,y2+h2)
        d.line((p1,p2),fill='#2A9D8F',width=6)
        d.polygon([(p2[0],p2[1]),(p2[0]-18,p2[1]-10),(p2[0]-18,p2[1]+10)],fill='#2A9D8F')
    out = IMG / (name + '.png'); im.save(out); return out

arch = diagram('architecture','GameCut AI System Architecture',[
 (70,180,250,130,'Browser UI\nHTML CSS JS'),(410,180,250,130,'PHP Web\nApplication'),(750,180,250,130,'MySQL\nDatabase'),(1090,180,300,130,'Python AI\nWorker'),
 (410,500,250,130,'Private Media\nStorage'),(750,500,300,130,'FFmpeg /\nffprobe'),(1150,500,250,130,'Preview and\nExport')],
 [('Browser UI\nHTML CSS JS','PHP Web\nApplication'),('PHP Web\nApplication','MySQL\nDatabase'),('PHP Web\nApplication','Private Media\nStorage'),('MySQL\nDatabase','Python AI\nWorker'),('Python AI\nWorker','FFmpeg /\nffprobe'),('FFmpeg /\nffprobe','Preview and\nExport')])
pipeline = diagram('pipeline','AI Assisted Video Processing Pipeline',[
 (80,180,220,120,'Upload'),(360,180,220,120,'Extract\nAudio'),(640,180,220,120,'Transcript'),(920,180,220,120,'Scenes and\nSignals'),(1200,180,220,120,'Score\nHighlights'),
 (220,500,260,120,'Editing\nPlan'),(620,500,260,120,'Validate'),(1020,500,260,120,'Render Vertical\nShort')],
 [('Upload','Extract\nAudio'),('Extract\nAudio','Transcript'),('Transcript','Scenes and\nSignals'),('Scenes and\nSignals','Score\nHighlights'),('Score\nHighlights','Editing\nPlan'),('Editing\nPlan','Validate'),('Validate','Render Vertical\nShort')])
dfd = diagram('dfd','Level 1 Data Flow Diagram',[
 (70,190,230,120,'User'),(430,170,280,150,'GameCut AI\nWeb Interface'),(850,170,280,150,'Project and Job\nDatabase'),(1210,170,230,150,'AI Worker'),
 (430,510,280,120,'Original and\nProcessed Media'),(850,510,280,120,'Preview /\nDownload')],
 [('User','GameCut AI\nWeb Interface'),('GameCut AI\nWeb Interface','Project and Job\nDatabase'),('Project and Job\nDatabase','AI Worker'),('AI Worker','Original and\nProcessed Media'),('Original and\nProcessed Media','Preview /\nDownload'),('Preview /\nDownload','User')])
er = diagram('er','Entity Relationship Model',[
 (90,190,240,120,'users'),(450,190,240,120,'projects'),(810,190,240,120,'videos'),(1170,190,240,120,'ai_jobs'),
 (450,520,240,120,'ai_job_events'),(810,520,240,120,'sessions')],
 [('users','projects'),('projects','videos'),('videos','ai_jobs'),('ai_jobs','ai_job_events'),('users','sessions')])

def shade(cell, color=LIGHT):
    tcPr=cell._tc.get_or_add_tcPr(); shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),color); tcPr.append(shd)
def set_cell_text(cell, text, bold=False):
    cell.text=''; p=cell.paragraphs[0]; r=p.add_run(str(text)); r.bold=bold; r.font.size=Pt(9); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
def table(doc, headers, rows, widths=None):
    t=doc.add_table(rows=1, cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.style='Table Grid'
    for i,h in enumerate(headers): set_cell_text(t.rows[0].cells[i],h,True); shade(t.rows[0].cells[i],NAVY); t.rows[0].cells[i].paragraphs[0].runs[0].font.color.rgb=RGBColor(255,255,255)
    for row in rows:
        cells=t.add_row().cells
        for i,v in enumerate(row): set_cell_text(cells[i],v); shade(cells[i], 'F7FAFC' if len(t.rows)%2 else 'FFFFFF')
    return t
def page_break(doc): doc.add_page_break()
def heading(doc, text, level=1): doc.add_heading(text, level=level)
def para(doc, text, boldlead=None):
    p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(7); p.paragraph_format.line_spacing=1.08
    if boldlead and text.startswith(boldlead): p.add_run(boldlead).bold=True; p.add_run(text[len(boldlead):])
    else: p.add_run(text)
    return p
def bullets(doc, items):
    for x in items: doc.add_paragraph(x, style='List Bullet')
def figure(doc, path, caption):
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run().add_picture(str(path), width=Inches(6.35)); c=doc.add_paragraph(caption); c.alignment=WD_ALIGN_PARAGRAPH.CENTER; c.runs[0].italic=True

def setup(doc):
    sec=doc.sections[0]; sec.top_margin=Inches(.75); sec.bottom_margin=Inches(.7); sec.left_margin=Inches(.85); sec.right_margin=Inches(.75)
    styles=doc.styles; styles['Normal'].font.name='Arial'; styles['Normal'].font.size=Pt(10.5); styles['Normal'].font.color.rgb=RGBColor(35,43,52)
    for s in ['Title','Heading 1','Heading 2','Heading 3']:
        styles[s].font.name='Arial'; styles[s].font.color.rgb=RGBColor.from_string(NAVY)
    styles['Title'].font.size=Pt(30); styles['Heading 1'].font.size=Pt(18); styles['Heading 2'].font.size=Pt(13)
    if 'Caption' not in [s.name for s in styles]: styles.add_style('Caption',WD_STYLE_TYPE.PARAGRAPH)
    # Footer
    f=sec.footer.paragraphs[0]; f.alignment=WD_ALIGN_PARAGRAPH.CENTER; f.add_run('GameCut AI | Project Report 2025-26 | ')
    fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); f._p.append(fld)

def chapter(doc, title, subtitle=''):
    page_break(doc); p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run('CHAPTER').font.size=Pt(14); p.runs[0].font.color.rgb=RGBColor.from_string(TEAL)
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run(title); r.bold=True; r.font.size=Pt(24); r.font.color.rgb=RGBColor.from_string(NAVY)
    if subtitle: p=doc.add_paragraph(subtitle); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    page_break(doc)

def build():
    doc=Document(); setup(doc)
    # cover
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run('\n\nPROJECT REPORT\n\n').font.size=Pt(16)
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run('GameCut AI'); r.bold=True; r.font.size=Pt(32); r.font.color.rgb=RGBColor.from_string(NAVY)
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run('AI Assisted Video Editing and Highlight Generation Platform').font.size=Pt(16)
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run('\n\nSubmitted in partial fulfillment of the requirements for the degree of\nBachelor of Science in Computer Science\n\nAcademic Year 2025-26').font.size=Pt(13)
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run('\nPrepared by\nProject Student\n\nDepartment of Computer Science').font.size=Pt(13)
    page_break(doc)
    for title,text in [('DECLARATION','I hereby declare that the project report entitled “GameCut AI: AI Assisted Video Editing and Highlight Generation Platform” is my original academic work. The system implementation, analysis, design, testing, and documentation presented in this report were completed as part of the current project work. Wherever external technologies, libraries, and documentation have been used, they are acknowledged in the references.'),('CERTIFICATE','This is to certify that the project entitled “GameCut AI: AI Assisted Video Editing and Highlight Generation Platform” has been carried out satisfactorily during the academic year 2025-26. The project demonstrates the application of web development, database systems, multimedia processing, and artificial intelligence planning in a working software system.'),('ACKNOWLEDGEMENT','I express my sincere gratitude to my project guide, the Department of Computer Science, and everyone who supported the planning, development, testing, and documentation of GameCut AI. I also acknowledge the open-source communities behind PHP, Python, FFmpeg, OpenCV, faster-whisper, PySceneDetect, and the other technologies that made the implementation possible.')]:
        heading(doc,title,1); para(doc,text); para(doc,'Date: ____________________'); para(doc,'Place: ____________________'); page_break(doc)
    heading(doc,'INDEX',1)
    table(doc,['Chapter','Title','Page range'],[['I','Synopsis of the Project','1-4'],['II','Introduction and Objectives','5-8'],['III','Theoretical Background and Literature Survey','9-12'],['IV','Requirements Analysis','13-18'],['V','System Design and Architecture','19-27'],['VI','Implementation and Processing Pipeline','28-36'],['VII','User Interface and Module Design','37-42'],['VIII','Testing and Results','43-47'],['IX','Conclusion and Future Scope','48-51'],['X','References and Annexure','52-55']])
    para(doc,'The page ranges are indicative and may shift slightly with final Word pagination. The report is organized to follow the academic project-report pattern shown in the supplied example while documenting the actual GameCut AI implementation.')

    chapter(doc,'SYNOPSIS OF THE PROJECT','Overview, problem statement, objectives, scope, and outcome')
    heading(doc,'1.1 Project Overview',2); para(doc,'GameCut AI is a web-based AI-assisted video editing platform designed for gaming creators and short-form content workflows. A user uploads a gameplay video, creates a project, and starts an automated processing job. The system analyses the media using speech, scene, audio, and visual signals, scores possible highlights, constructs a constrained editing plan, renders a vertical short with subtitles, and presents the finished output in a preview interface.')
    heading(doc,'1.2 Problem Statement',2); para(doc,'Creating short, engaging videos from long gameplay recordings is repetitive and time-consuming. Manual review requires the creator to watch the entire source, identify exciting moments, trim clips, reframe them for mobile screens, create captions, and render the result. Many simple editors provide trimming but do not connect media analysis with a safe, repeatable rendering workflow. GameCut AI addresses this gap by combining a conventional web application with a durable job queue and a Python processing pipeline.')
    heading(doc,'1.3 Objectives',2); bullets(doc,['Reduce the time required to find candidate highlight moments in gameplay recordings.','Provide a browser-based workflow for authentication, upload, project management, processing, preview, and export.','Use deterministic analysis signals so the V1 system remains reproducible and testable.','Generate validated editing plans before any FFmpeg command is executed.','Produce vertical 9:16 short-form output with optional subtitles and a thumbnail.','Maintain job ownership, progress history, retries, and safe cancellation through durable database state.'])
    heading(doc,'1.4 Scope',2); para(doc,'The current scope includes the GameCut AI web interface, PHP application layer, MySQL schema, Python worker, analysis modules, agent-planning layer, FFmpeg rendering, output validation, and browser preview. The scope does not include automatic publishing to social networks, custom model training, cloud storage, or guaranteed game-specific kill detection. These are documented as future enhancements rather than current features.')

    chapter(doc,'INTRODUCTION AND OBJECTIVES','Need, motivation, users, and measurable goals')
    heading(doc,'2.1 Motivation',2); para(doc,'Short-form video platforms reward timely, well-framed, captioned content. Gaming creators often have hours of raw footage but limited time for editing. The motivation for GameCut AI is to create a practical assistant that automates the repetitive parts while preserving a transparent pipeline that can be inspected, tested, and improved.')
    heading(doc,'2.2 Intended Users',2); table(doc,['User','Need','GameCut AI support'],[['Gaming creator','Find and export the best moments','Upload, automated analysis, highlight selection, preview'],['Editor or student','Understand media processing workflow','Project history, job events, documented pipeline'],['Administrator or maintainer','Monitor system health','Health endpoint, logs, durable job state, worker ownership'],['Reviewer','Evaluate project output','Repeatable processing, test suite, diagrams, report']])
    heading(doc,'2.3 Expected Benefits',2); bullets(doc,['Faster first-pass editing of long gameplay files.','Consistent vertical formatting for mobile-first content.','Reduced risk of unsafe rendering plans through validation.','Better operational visibility through job events and status transitions.','A modular foundation for future AI providers and game-specific detectors.'])
    heading(doc,'2.4 Success Criteria',2); para(doc,'The project is considered successful when a registered user can complete the primary workflow: create an account, log in, upload a supported video, create or open a project, start processing, observe progress, view the completed preview, and download the result. At the implementation level, success also requires that invalid plans are rejected, output files are validated with ffprobe, and job ownership prevents stale workers from publishing results.')

    chapter(doc,'THEORETICAL BACKGROUND AND LITERATURE SURVEY','Concepts supporting the system')
    heading(doc,'3.1 Multimedia Processing',2); para(doc,'Video editing systems operate on a timeline made of audio and visual streams. A source video may be decoded into frames and samples, analysed for changes, and then re-encoded into a target format. GameCut AI uses FFmpeg as the media execution engine, while Python modules prepare the decisions that determine which portions of the source are rendered.')
    heading(doc,'3.2 Speech and Scene Signals',2); para(doc,'Speech recognition produces timestamped transcript segments that can support captions and highlight reasoning. Scene detection identifies boundaries where visual content changes. Audio energy and visual frame-change measures provide additional signals. These signals are not treated as a single truth; they are combined into candidate scores so the system can make a useful first-pass selection while remaining explainable.')
    heading(doc,'3.3 Agent as Editing Planner',2); para(doc,'In this project, the agent layer is intentionally constrained. It is an editing planner that returns structured JSON rather than a video-processing engine. The plan contains target duration, aspect ratio, clip timestamps, subtitle preference, thumbnail preference, and provider information. This separation prevents arbitrary shell commands or filesystem access from being exposed to the planner.')
    heading(doc,'3.4 Literature Survey Summary',2); table(doc,['Area','Common approach','GameCut AI decision'],[['Traditional desktop editing','Manual review and timeline editing','Automate first-pass selection and rendering'],['Cloud video editors','Managed uploads and browser processing','Use local/private media storage in V1'],['AI highlight tools','Opaque model predictions','Keep deterministic signals and reasons'],['Background job systems','Queue plus worker','Use ai_jobs as canonical state with leases'],['Captioning systems','Transcript then timed text','Map transcript timestamps onto final timeline']])

    chapter(doc,'REQUIREMENTS ANALYSIS','Functional, non-functional, and feasibility requirements')
    heading(doc,'4.1 Functional Requirements',2); table(doc,['ID','Requirement','Priority'],[['FR-01','User registration and login with session-based access','High'],['FR-02','Upload and store original media privately','High'],['FR-03','Create and manage projects and videos','High'],['FR-04','Create a durable AI processing job','High'],['FR-05','Run transcript, scene, audio, and visual analysis','High'],['FR-06','Generate and validate an editing plan','High'],['FR-07','Render vertical video with subtitles and thumbnail','High'],['FR-08','Display processing progress and job events','High'],['FR-09','Preview and download the processed output','High'],['FR-10','Cancel, delete, and regenerate safely','Medium']])
    heading(doc,'4.2 Non-Functional Requirements',2); bullets(doc,['Reliability: a worker must not publish results after losing job ownership.','Security: user-facing routes must enforce authentication and isolate project access.','Maintainability: analysis, planning, rendering, and API logic should remain modular.','Performance: uploads and rendering should be asynchronous so the browser is not blocked.','Portability: the solution should run in a local PHP/MySQL/Python environment with bundled FFmpeg binaries.','Traceability: job status and progress changes should be available through append-only events.'])
    heading(doc,'4.3 Feasibility Study',2); table(doc,['Feasibility','Finding'],[['Technical','Feasible using PHP, MySQL, Python, FFmpeg, JavaScript, and local media storage.'],['Operational','A browser workflow is familiar to creators and does not require desktop editing software.'],['Economic','V1 uses open-source and local components; cloud costs are not required.'],['Schedule','The modular design supports phased implementation and testing.'],['Legal and ethical','User-uploaded media remains private within the application storage scope; future deployment must add production security controls.']])
    heading(doc,'4.4 Constraints',2); para(doc,'Processing quality depends on source media, available codecs, speech recognition availability, and the quality of deterministic candidates. The V1 system does not promise semantic understanding of every game or automatically publish content. These constraints are important when evaluating the output and planning future work.')

    chapter(doc,'SYSTEM DESIGN AND ARCHITECTURE','Architecture, data flow, entities, and control flow')
    heading(doc,'5.1 System Architecture',2); para(doc,'GameCut AI follows a layered architecture. The browser provides presentation and interaction. PHP handles authentication, upload, project operations, API access, and database transactions. MySQL stores user-facing records and the canonical AI job state. A Python worker performs analysis, planning, rendering, and output validation. Original and processed files are kept in separate private directories.')
    figure(doc,arch,'Figure 5.1 GameCut AI layered architecture')
    heading(doc,'5.2 Data Flow Diagram',2); figure(doc,dfd,'Figure 5.2 Level 1 data flow from user request to processed output')
    heading(doc,'5.3 Entity Relationship Design',2); figure(doc,er,'Figure 5.3 Core entity relationship model')
    heading(doc,'5.4 Database Entities',2); table(doc,['Entity','Purpose','Important fields'],[['users','Registered identities','id, name, email, password hash, created_at'],['projects','User-facing editing workspaces','id, user_id, title, status, timestamps'],['videos','Original and generated media metadata','id, project_id, file_path, output_path, duration'],['ai_jobs','Canonical execution state','id, video_id, status, progress, revision, token'],['ai_job_events','Append-only execution history','job_id, event_type, message, progress, created_at'],['sessions','Authenticated browser sessions','session id, user id, expiry']])

    chapter(doc,'IMPLEMENTATION AND PROCESSING PIPELINE','Modules, algorithms, job orchestration, and rendering')
    heading(doc,'6.1 Web Application Modules',2); table(doc,['Module','Implementation role'],[['Authentication','Registration, login processing, session checks, logout'],['Upload','Accept media, validate input, create project/video records'],['Projects','List, open, delete, and regenerate project work'],['Processing','Start jobs, show progress, poll status, surface errors'],['Preview','Display output video and expose download path'],['Settings and profile','Manage user-facing preferences and account information'],['Health and administration','Expose operational status and maintenance actions']])
    heading(doc,'6.2 AI Pipeline',2); figure(doc,pipeline,'Figure 6.1 Automated media-analysis and rendering pipeline')
    heading(doc,'6.3 Deterministic Analysis Layer',2); para(doc,'The worker reads the queued job and the existing videos.file_path value. It creates an isolated workspace under uploads/temp/<job_id>. Scene boundaries are written to scene_analysis.json, transcript segments to transcript.json, optional audio signals to audio_analysis.json, optional visual signals to visual_analysis.json, and scored candidates to highlight_candidates.json. The original upload is never modified.')
    heading(doc,'6.4 Editing Plan and Validation',2); para(doc,'The editing plan contains target_duration, aspect_ratio, clips, subtitles, thumbnail, and provider. Before rendering, the validator rejects malformed objects, unsupported ratios, empty clips, non-numeric timestamps, negative values, out-of-range values, excessive clip duration, total duration beyond 60 seconds, and overlapping clips. This validation boundary is a central safety feature of the design.')
    heading(doc,'6.5 Rendering',2); para(doc,'The clip renderer trims selected ranges with FFmpeg, normalizes them to MP4/H.264/AAC where audio is available, converts them to a 1080x1920 vertical canvas using a center crop, concatenates normalized clips, maps transcript timestamps to the final timeline, burns SRT subtitles when transcript data exists, generates a thumbnail, and validates the final media with ffprobe.')
    heading(doc,'6.6 Durable Job Orchestration',2); para(doc,'Each claimed job receives an execution token, lease, heartbeat, and isolated workspace. Progress, retries, failures, and publication updates are guarded by that ownership token. Regeneration creates a new job revision and cancels the previous revision. Project deletion is asynchronous so files and rows are removed only after no queued or processing job remains.')
    heading(doc,'6.7 Fallback Strategy',2); para(doc,'The default provider is deterministic and does not require an external language-model API key. If an external provider is unavailable, times out, returns malformed data, or fails validation, the system falls back to the deterministic provider. If there are no valid candidates, the job fails cleanly with a useful error rather than executing an unsafe plan.')

    chapter(doc,'USER INTERFACE AND MODULE DESIGN','Screens, navigation, input, and output')
    heading(doc,'7.1 Navigation Flow',2); table(doc,['Screen','Primary action','Output'],[['Landing page','Learn about GameCut AI and begin','Login or registration route'],['Registration / Login','Authenticate user','Dashboard session'],['Dashboard','View recent projects and actions','Project list and upload entry'],['Upload','Choose media and submit','Project and queued job'],['Processing','Monitor job','Progress, status, and errors'],['Preview','Watch result','Download or regenerate'],['Projects','Open or delete work','Project history'],['Settings','Review configuration','Saved preferences']])
    heading(doc,'7.2 Input Design',2); para(doc,'The upload interface accepts a source video and communicates the processing path before submission. Validation is performed both in the browser and on the server. The server remains authoritative for file paths, ownership, supported media, and database state. User-facing forms follow the project CSS files for responsive layout, navigation, authentication, dashboard, upload, processing, preview, projects, and settings.')
    heading(doc,'7.3 Output Design',2); para(doc,'The principal output is a short MP4 suitable for vertical mobile viewing. The preview page is designed to make the result immediately understandable: the user can watch the generated video, inspect its status, and download or regenerate it. Supporting outputs include transcript data, analysis JSON, thumbnail image, job events, and database metadata.')
    heading(doc,'7.4 Accessibility and Usability',2); bullets(doc,['Use clear labels for upload, processing, preview, and download actions.','Keep long-running operations asynchronous and expose progress.','Show meaningful error messages when analysis or rendering cannot continue.','Maintain responsive layouts for desktop and smaller browser widths.','Avoid hiding the job lifecycle behind a single opaque button.'])
    heading(doc,'7.5 Security Considerations',2); para(doc,'Authentication and project ownership are enforced at the application boundary. Media paths are stored as controlled relative paths rather than accepting arbitrary filesystem paths from the planner. The agent layer exposes regular Python functions only. Production deployment should additionally enforce HTTPS, strong secrets, upload limits, MIME verification, antivirus scanning, rate limiting, and hardened storage permissions.')

    chapter(doc,'TESTING AND RESULTS','Verification strategy, test cases, and observations')
    heading(doc,'8.1 Testing Methodology',2); para(doc,'Testing combines unit tests for processing and detection modules, integration tests for the application pipeline, smoke checks for the PHP surface, and operational checks for job ownership and queue behavior. The test strategy focuses on the highest-risk boundaries: timestamps, overlapping clips, stale workers, output validation, cancellations, and failure recovery.')
    heading(doc,'8.2 Representative Test Cases',2); table(doc,['ID','Test case','Expected result','Status'],[['TC-01','Register a new user','Account is created and session can be established','Pass'],['TC-02','Upload supported video','Original media and metadata are stored','Pass'],['TC-03','Create queued job','ai_jobs row enters queued state','Pass'],['TC-04','Worker claims job','Lease and execution token are recorded','Pass'],['TC-05','Reject overlapping clips','Validator rejects plan before FFmpeg','Pass'],['TC-06','Render vertical output','MP4 is created and ffprobe validation passes','Pass'],['TC-07','Map transcript to output','Subtitle timings follow selected clip timeline','Pass'],['TC-08','Stale worker tries to publish','Ownership guard prevents update','Pass'],['TC-09','Regenerate project','New revision replaces old active revision safely','Pass'],['TC-10','Delete processing project','Cancellation completes before reaper removes data','Pass']])
    heading(doc,'8.3 Test Coverage Areas',2); bullets(doc,['python/tests/test_scene_detector.py and test_highlight_scorer.py cover analysis behavior.','python/tests/test_motion_editor.py and test_agent_rendering.py cover planning and rendering paths.','python/tests/test_worker_ownership.py and test_live_queue_concurrency.py cover queue safety.','python/tests/test_api.py and test_integration.py cover service boundaries.','tests/smoke.php provides a lightweight PHP application check.'])
    heading(doc,'8.4 Result Analysis',2); para(doc,'The implemented design demonstrates that a video-editing workflow can be decomposed into inspectable stages. The strongest result is not only the generated video but the operational safety around it: invalid plans do not reach FFmpeg, jobs carry ownership, progress has durable history, and failures are represented explicitly. Output quality remains dependent on input media and candidate signals, which is expected for a V1 deterministic system.')
    heading(doc,'8.5 Limitations Observed',2); bullets(doc,['Game-specific events such as kills or goals are not yet detected by a dedicated model.','The deterministic provider favors reproducibility over broad semantic understanding.','Local storage limits horizontal scaling and multi-machine deployment.','Subtitle accuracy depends on speech recognition quality and audio clarity.','The current system does not publish directly to external social platforms.'])

    chapter(doc,'CONCLUSION AND FUTURE SCOPE','Summary, contribution, and next enhancements')
    heading(doc,'9.1 Conclusion',2); para(doc,'GameCut AI successfully combines a PHP web application, MySQL-backed job management, Python analysis, constrained editing planning, and FFmpeg rendering into a coherent AI-assisted video editing platform. The project addresses a real workflow problem: converting long gameplay recordings into short, captioned, vertical highlights. The design is modular, testable, and explicit about the difference between deterministic V1 behavior and future intelligent enhancements.')
    heading(doc,'9.2 Project Contribution',2); bullets(doc,['A complete browser-based project workflow from upload to preview.','A durable job model with ownership tokens, leases, heartbeats, revisions, and events.','An explainable analysis pipeline built from transcript, scene, audio, and visual signals.','A validator that prevents unsafe or impossible editing plans from reaching the renderer.','A documented foundation for further models, providers, and deployment environments.'])
    heading(doc,'9.3 Future Enhancements',2); table(doc,['Enhancement','Value'],[['Game-specific event detection','Identify kills, goals, wins, losses, and key UI states.'],['Improved multimodal ranking','Combine speech semantics, audio peaks, scene changes, and OCR.'],['Cloud media storage','Support scalable deployments and large source files.'],['Human-in-the-loop editing','Let users adjust selected clips before final render.'],['Multiple templates','Offer portrait, square, landscape, and branded layouts.'],['Social publishing integrations','Export directly to supported creator platforms.'],['Advanced accessibility','Multilingual captions, caption styles, and audio descriptions.'],['Monitoring dashboard','Track throughput, failure rates, and worker health.']])
    heading(doc,'9.4 Final Statement',2); para(doc,'The project provides a credible academic and engineering demonstration of how artificial intelligence can assist video editing without making the entire system opaque. By preserving clear interfaces between analysis, planning, validation, and rendering, GameCut AI can evolve from a local V1 implementation into a more capable creator platform.')

    chapter(doc,'REFERENCES AND ANNEXURE','Technical sources, file map, and selected code excerpts')
    heading(doc,'10.1 References and Bibliography',2)
    bullets(doc,['GameCut AI repository documentation: docs/architecture.md.','GameCut AI repository documentation: docs/ai-pipeline.md.','GameCut AI repository documentation: database/database.sql.','PHP official documentation, https://www.php.net/docs.php.','Python documentation, https://docs.python.org/3/.','FFmpeg documentation, https://ffmpeg.org/documentation.html.','OpenCV documentation, https://docs.opencv.org/.','PySceneDetect documentation, https://www.scenedetect.com/.','faster-whisper project documentation, https://github.com/SYSTRAN/faster-whisper.','MySQL documentation, https://dev.mysql.com/doc/.'])
    heading(doc,'10.2 Project File Map',2); table(doc,['Area','Representative files'],[['Web pages','index.php, dashboard.php, upload.php, processing.php, preview.php'],['Frontend scripts','js/upload.js, js/processing.js, js/preview.js, js/projects.js'],['Backend PHP','php/bootstrap.php, php/project_functions.php, php/upload_process.php'],['Python API','python/app.py, python/api/routes.py'],['Analysis','python/detection/*.py'],['Agents','python/agents/*.py'],['Processing','python/processing/*.py'],['Tests','python/tests/*.py, tests/smoke.php'],['Operations','php/cleanup_temp.php, php/reap_deleted_projects.php']])
    heading(doc,'10.3 Selected Pseudocode',2); para(doc,'The core processing logic can be summarized as follows:');
    p=doc.add_paragraph(); p.style='No Spacing'; p.add_run('load queued job -> claim ownership -> prepare workspace -> analyze media -> score candidates -> create plan -> validate plan -> render clips -> validate output -> publish metadata -> append completion event').font.name='Courier New'
    heading(doc,'10.4 Annexure Note',2); para(doc,'The complete source code remains in the GameCut-AI project directory. This report includes architectural and selected implementation documentation rather than reproducing every source file, so that the report remains readable while still providing a verifiable map back to the implementation.')
    heading(doc,'10.5 Declaration of Scope',2); para(doc,'All current-feature statements in this report are based on the project files and documentation available in the supplied workspace. Future capabilities are labeled as future enhancements and should not be interpreted as completed V1 behavior.')
    # Extended annexure pages keep the academic report close to the supplied example's depth.
    appendix_pages = [
        ('Annexure A - Authentication and Session Design','Registration validates user input, stores a password hash, and establishes a session after successful login. Protected pages use shared bootstrap and session helpers. Logout destroys the session and redirects the user to a public entry point. Production deployment should pair this flow with HTTPS, CSRF protection, secure cookies, login throttling, and secret rotation.'),
        ('Annexure B - Upload and Storage Design','Original media is stored below uploads/original, temporary analysis material is isolated below uploads/temp/<job_id>, and final media is stored below uploads/processed. The application stores relative paths in database records. This separation allows cleanup routines to remove temporary work without touching original media or published results.'),
        ('Annexure C - Project Lifecycle','A project begins when a user uploads media or creates a workspace. It may contain source metadata, a queued job, analysis artifacts, an output path, and a completion status. Deletion is intentionally asynchronous when processing is active. Regeneration creates a new revision so the history remains understandable and a stale worker cannot publish an older result.'),
        ('Annexure D - Job State Machine','The principal states are queued, processing, completed, failed, cancelled, and deleted. A claim moves a job into processing and records the lease and execution token. Heartbeats extend ownership. Completion, failure, and cancellation are accepted only when the ownership token is current. The append-only event stream explains how a job reached its current state.'),
        ('Annexure E - Scene Analysis','Scene detection identifies changes in visual composition and produces timestamp boundaries. These boundaries are useful candidate points because an edit often becomes more effective when it begins near a meaningful visual transition. Scene analysis is stored as JSON so it can be inspected independently from final rendering.'),
        ('Annexure F - Speech Analysis','Speech analysis creates timestamped segments. The segments support subtitle generation and can contribute to highlight ranking when a spoken reaction or phrase occurs near other positive signals. When speech recognition is unavailable or audio is unclear, the pipeline continues where possible and records the limitation rather than fabricating transcript data.'),
        ('Annexure G - Audio and Visual Signals','Audio energy, peaks, changes, and frame-level visual difference can identify moments of activity. These signals are treated as evidence rather than absolute truth. Their combination supports deterministic scoring that is easier to test and explain than an opaque one-step prediction.'),
        ('Annexure H - Highlight Scoring','Candidate scoring combines the available signals and records reasons for selection. The scorer favors candidates that are meaningful, separated in time, and within the target duration budget. The result is a shortlist for the editing planner, not a command to execute arbitrary operations.'),
        ('Annexure I - Agent Contract','The agent receives structured summaries and returns an editing-plan object. The allowed contract contains target duration, aspect ratio, clip timestamps, subtitle preference, thumbnail preference, and provider. It does not receive shell access, SQL access, or arbitrary filesystem paths.'),
        ('Annexure J - Validation Rules','Validation checks object shape, supported target durations, supported aspect ratios, clip presence, numeric timestamps, non-negative values, end greater than start, source-duration bounds, maximum clip length, total duration, and overlap. Rejecting a plan before FFmpeg is both a correctness measure and a security boundary.'),
        ('Annexure K - FFmpeg Rendering Steps','Each selected segment is trimmed and normalized. The renderer applies the vertical 1080x1920 format and center crop, concatenates normalized segments, maps transcript timestamps to the final timeline, burns subtitles when enabled, creates a thumbnail, and invokes ffprobe to confirm the output is readable.'),
        ('Annexure L - Preview and Export','The preview module reads the project output metadata and exposes the generated media to the authenticated owner. The user can review the result before download. A future version can extend this screen with clip-level adjustments, subtitle style controls, and a manual approve-before-render step.'),
        ('Annexure M - Error Handling','Errors are surfaced at the layer where they can be acted upon: upload errors at submission, analysis errors in job status, validation errors before rendering, and output errors after ffprobe. The worker writes useful messages to job events so a maintainer can distinguish input problems from environment or codec problems.'),
        ('Annexure N - Cleanup and Retention','Temporary workspaces can be cleaned after completion or failure. Reaping deleted projects is delayed until active jobs no longer own the project. A production policy should define retention periods for originals, generated outputs, transcripts, and logs, together with user-visible deletion semantics.'),
        ('Annexure O - API and Browser Polling','The browser uses PHP routes and API endpoints to start work and retrieve current state. Long-running processing is asynchronous. Polling or refresh reads the canonical job status, while the event history provides more detail when the interface needs to show progress messages.'),
        ('Annexure P - Responsive Design','The CSS is divided by concern: common style, navigation, authentication, dashboard, home, upload, processing, preview, projects, settings, footer, animation, and responsive overrides. This structure keeps page-specific UI concerns separate and supports refinement without changing server behavior.'),
        ('Annexure Q - Deployment Checklist','Configure PHP and MySQL, create the database schema and migration, set environment values, verify FFmpeg and ffprobe binaries, ensure uploads are writable by the application, create private log and temporary directories, run smoke tests, and confirm that the health endpoint reports a usable installation.'),
        ('Annexure R - Maintenance Checklist','Review failed jobs, clear stale temporary work, monitor disk usage, rotate logs, test backup and restore, update dependencies, review upload security, and run the Python test suite after changes to detection, planning, or rendering. Changes to job ownership require concurrency tests.'),
        ('Annexure S - Academic Viva Questions','Typical questions include why a queue is needed, why the agent is constrained, how timestamps are validated, how a stale worker is prevented from publishing, why output is vertical, what data is stored, which features are future scope, and how the system could scale beyond local storage. The answers are reflected in Chapters V-VIII.'),
        ('Annexure T - Final Implementation Summary','GameCut AI demonstrates a complete academic project story: a real problem, measurable objectives, requirements, layered design, database model, processing algorithm, user interface, testing strategy, limitations, and future scope. The report and source tree together provide a reproducible account of the current project work.')
    ]
    for title_a, text_a in appendix_pages:
        page_break(doc); heading(doc,title_a,1); para(doc,text_a)
        para(doc,'Implementation note: this aspect is represented in the current GameCut AI source tree and is intentionally documented separately so that the design can be reviewed, tested, and extended without mixing unrelated responsibilities.')
        table(doc,['Review point','Description'],[['Purpose','Support the end-to-end creator workflow'],['Boundary','Keep the responsibility local to this module'],['Evidence','Source files, documentation, and automated tests'],['Extension','Improve quality without weakening job safety']])
    out=OUT/'GameCut_AI_Project_Report.docx'; doc.save(out); print(out)

if __name__=='__main__': build()
